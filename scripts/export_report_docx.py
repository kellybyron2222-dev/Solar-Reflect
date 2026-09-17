"""Export report.md to a Word .docx for drafting in Word."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "report.md"
OUT = ROOT / "report.docx"


def latex_to_plain(s: str) -> str:
    t = s
    t = t.replace(r"\mathrm{km}", "km")
    t = t.replace(r"\mathrm{m}", "m")
    t = t.replace(r"\mathrm{min}", "min")
    t = t.replace(r"\mathrm{s}", "s")
    t = t.replace(r"\mathrm{W}", "W")
    t = t.replace(r"\mathrm{J}", "J")
    t = t.replace(r"\mathrm{MJ}", "MJ")
    t = t.replace(r"\mathrm{kJ}", "kJ")
    t = t.replace(r"\,\mathrm{km}^2", " km²")
    t = t.replace(r"\,\mathrm{m^2}", " m²")
    t = t.replace(r"\mathrm{km}^2", "km²")
    t = t.replace(r"\mathrm{m^2}", "m²")
    t = t.replace(r"\text{image}", "image")
    t = t.replace(r"\text{useful}", "useful")
    t = t.replace(r"\text{ground}", "ground")
    t = t.replace(r"\text{in}", "in")
    t = t.replace(r"\text{one}", "one")
    t = t.replace(r"\text{horizon}", "horizon")
    t = t.replace(r"\approx", "≈")
    t = t.replace(r"\times", "×")
    t = t.replace(r"\cdot", "·")
    t = t.replace(r"\sim", "~")
    t = t.replace(r"\alpha", "α")
    t = t.replace(r"\eta", "η")
    t = t.replace(r"\gamma", "γ")
    t = t.replace(r"\varepsilon", "ε")
    t = t.replace(r"\mu", "μ")
    t = t.replace(r"\pi", "π")
    t = t.replace(r"\rho", "ρ")
    t = t.replace(r"\tau", "τ")
    t = t.replace(r"\Delta", "Δ")
    t = t.replace(r"\min", "min")
    t = t.replace(r"\bigl", "")
    t = t.replace(r"\bigr", "")
    t = t.replace(r"\left", "")
    t = t.replace(r"\right", "")
    t = t.replace(r"\frac", "")
    t = t.replace(r"\,", " ")
    t = t.replace(r"\ ", " ")
    t = t.replace("{", "").replace("}", "")
    t = re.sub(r"\$\$", "", t)
    t = re.sub(r"\\\((.*?)\\\)", r"\1", t)
    t = re.sub(r"\\\[(.*?)\\\]", r"\1", t, flags=re.S)
    t = t.replace("$", "")
    t = t.replace("  ", " ")
    return t.strip()


def add_runs(paragraph, text: str) -> None:
    text = latex_to_plain(text)
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            paragraph.add_run(part)


def set_run_font(run, name="Calibri", size=11) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)


def style_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.15)
    section.right_margin = Inches(1.15)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(0x1C, 0x19, 0x16)
    pf = style.paragraph_format
    pf.space_after = Pt(10)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    for i, size in ((1, 22), (2, 16), (3, 13)):
        hs = doc.styles[f"Heading {i}"]
        hs.font.name = "Calibri"
        hs.font.size = Pt(size)
        hs.font.bold = True
        hs.font.color.rgb = RGBColor(0x1C, 0x19, 0x16)
        hs.paragraph_format.space_before = Pt(16 if i > 1 else 0)
        hs.paragraph_format.space_after = Pt(8)


def image_path(md_path: str) -> Path | None:
    p = ROOT / md_path.replace("\\", "/")
    if p.suffix.lower() == ".svg":
        png = p.with_suffix(".png")
        if png.exists():
            return png
    if p.exists() and p.suffix.lower() in {".png", ".jpg", ".jpeg"}:
        return p
    return None


def export() -> Path:
    md = SRC.read_text(encoding="utf-8")
    doc = Document()
    style_doc(doc)

    lines = md.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        raw = line.rstrip()
        if not raw:
            i += 1
            continue
        if raw.strip() == "---":
            i += 1
            continue
        if raw.startswith("# "):
            doc.add_heading(latex_to_plain(raw[2:]), level=1)
            i += 1
            continue
        if raw.startswith("## "):
            doc.add_heading(latex_to_plain(raw[3:]), level=2)
            i += 1
            continue
        if raw.startswith("### "):
            doc.add_heading(latex_to_plain(raw[4:]), level=3)
            i += 1
            continue
        img = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", raw)
        if img:
            caption, rel = img.group(1), img.group(2)
            pic = image_path(rel)
            if pic:
                doc.add_picture(str(pic), width=Inches(6.2))
                last = doc.paragraphs[-1]
                last.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = cap.add_run(latex_to_plain(caption))
            run.italic = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x6A, 0x63, 0x5B)
            i += 1
            continue
        if raw.startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*:?--", lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [latex_to_plain(c.strip()) for c in lines[i].strip().strip("|").split("|")]
                if not re.match(r"^:?-+:?$", cells[0].replace(" ", "")):
                    rows.append(cells)
                i += 1
            if rows:
                cols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=cols)
                table.style = "Table Grid"
                for r_i, row in enumerate(rows):
                    for c_i in range(cols):
                        cell = table.cell(r_i, c_i)
                        cell.text = row[c_i] if c_i < len(row) else ""
                        for p in cell.paragraphs:
                            for run in p.runs:
                                run.font.size = Pt(10)
                                if r_i == 0:
                                    run.bold = True
            continue
        if raw.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            add_runs(p, raw[2:])
            for run in p.runs:
                run.italic = True
            i += 1
            continue
        if raw.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, raw[2:])
            i += 1
            continue
        if raw.startswith(r"\[") or raw.startswith("$$"):
            block = [raw]
            i += 1
            while i < len(lines) and not (lines[i].strip().startswith(r"\]") or lines[i].strip() == "$$"):
                block.append(lines[i])
                i += 1
            if i < len(lines):
                block.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(latex_to_plain(" ".join(block)))
            run.italic = True
            continue
        para_lines = [raw]
        i += 1
        while i < len(lines):
            nxt = lines[i].rstrip()
            if not nxt or nxt.startswith(("#", "|", "!", "-", ">", "---", r"\[", "$$")):
                break
            para_lines.append(nxt)
            i += 1
        p = doc.add_paragraph()
        add_runs(p, " ".join(para_lines))

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = export()
    print(path)
